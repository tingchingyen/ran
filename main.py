from pptx import Presentation
from glob import glob
import streamlit as st
st.write(glob('*'))



# Open an existing PowerPoint presentation
# prs = Presentation("template.pptx")

# Get the first slide of the presentation
# slide = prs.slides[0]

# Add a text box to the slide
# textbox = slide.shapes.add_textbox(left=50, top=50, width=200, height=100)

# Set the text of the text box
# textbox.text_frame.text = "Hello, world!"


# prs.save("presentation.pptx")
